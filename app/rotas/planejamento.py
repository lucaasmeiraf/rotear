"""Planejamento: nada aqui é gravado, exceto a linha agregada de metricas_uso (sem endereços)."""
import io
import threading
import time
from datetime import datetime
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, field_validator
from sqlalchemy.orm import Session

from app.db import get_db
from app.deps import exigir_login
from app.entrada import MAX_BYTES, ArquivoInvalido, ler_planilha
from app.geocodificacao import dentro_da_regiao
from app.models import MetricaUso, Usuario
from app.roteamento import Ponto
from app.roteamento.links import trechos
from app.roteamento.matriz import MatrizIndisponivel, obter_provedor
from app.roteamento.solver import MOTIVO_LIMITE, resolver, tempo_para
from app.rotas.admin import obter_configuracao

router = APIRouter()
FUSO = ZoneInfo("America/Sao_Paulo")
MAX_MOTORISTAS = 20
# OR-Tools usa 1 núcleo por cálculo; a VPS tem 2 vCPUs
_calculos = threading.BoundedSemaphore(2)


@router.get("/api/planejamento/config")
def config_planejamento(_: Usuario = Depends(exigir_login), db: Session = Depends(get_db)):
    c = obter_configuracao(db)
    partida = None
    if c.modo_partida == "endereco_fixo" and c.partida_lat is not None:
        partida = {"lat": c.partida_lat, "lng": c.partida_lng, "endereco": c.partida_endereco}
    return {"modo_partida": c.modo_partida, "partida": partida, "max_paradas": c.max_paradas,
            "solver_tempo_limite_s": c.solver_tempo_limite_s, "geocodificacao_ativa": c.geocodificacao_ativa}


@router.post("/api/entrada/arquivo")
async def enviar_arquivo(arquivo: UploadFile = File(...), _: Usuario = Depends(exigir_login)):
    conteudo = await arquivo.read(MAX_BYTES + 1)
    try:
        return ler_planilha(arquivo.filename or "", conteudo)
    except ArquivoInvalido as exc:
        raise HTTPException(422, str(exc)) from None


# --- Cálculo ------------------------------------------------------------------


class PontoIn(BaseModel):
    lat: float = Field(ge=-34, le=6)
    lng: float = Field(ge=-74, le=-34)


class ParadaIn(BaseModel):
    """Sem lat/lng (endereço não localizado), a parada entra no fim de uma rota e vai para o Google Maps pelo texto."""
    lat: float | None = Field(default=None, ge=-34, le=6)
    lng: float | None = Field(default=None, ge=-74, le=-34)
    endereco: str = Field(min_length=1, max_length=400)


class MotoristaIn(BaseModel):
    nome: str = Field(default="", max_length=60)


class CalculoIn(BaseModel):
    partida: PontoIn | None = None
    paradas: list[ParadaIn] = Field(min_length=1)
    motoristas: list[MotoristaIn] = Field(min_length=1, max_length=MAX_MOTORISTAS)
    max_por_motorista: int | None = Field(default=None, ge=1)
    fim: Literal["ultima_entrega", "partida"] = "ultima_entrega"

    @field_validator("motoristas")
    @classmethod
    def _nomes(cls, v):
        for i, m in enumerate(v):
            m.nome = m.nome.strip() or f"Motorista {i + 1}"
        return v


@router.post("/api/calcular")
def calcular(dados: CalculoIn, usuario: Usuario = Depends(exigir_login), db: Session = Depends(get_db)):
    cfg = obter_configuracao(db)
    if len(dados.paradas) > cfg.max_paradas:
        raise HTTPException(422, f"No máximo {cfg.max_paradas} paradas por planejamento.")

    # No modo endereço fixo, o ponto vem sempre da configuração (o operador não altera)
    if cfg.modo_partida == "endereco_fixo":
        if cfg.partida_lat is None:
            raise HTTPException(409, "O endereço de partida ainda não foi configurado pelo administrador.")
        partida = Ponto(cfg.partida_lat, cfg.partida_lng, cfg.partida_endereco)
    else:
        if dados.partida is None:
            raise HTTPException(422, "Informe a sua localização atual para usar como ponto de partida.")
        partida = Ponto(dados.partida.lat, dados.partida.lng)
    if not dentro_da_regiao(partida.lat, partida.lng):
        raise HTTPException(422, "O ponto de partida está fora da área atendida pelas rotas (DF e entorno).")

    paradas = [Ponto(p.lat, p.lng, p.endereco) for p in dados.paradas]
    # Nenhuma parada fica de fora: as que não têm posição confiável no mapa (não localizadas, fora da área ou sem rota
    # pelas ruas) não entram na otimização e vão para o fim da rota mais curta; o link leva o texto ao Google Maps.
    com_posicao = [i for i, p in enumerate(paradas)
                   if p.lat is not None and p.lng is not None and dentro_da_regiao(p.lat, p.lng)]
    sem_posicao = sorted(set(range(len(paradas))) - set(com_posicao))
    retorno = dados.fim == "partida"
    provedor = obter_provedor()
    n_mot, limite = len(dados.motoristas), dados.max_por_motorista

    inicio = time.monotonic()
    with _calculos:
        rotas: list[list[int]] = [[] for _ in range(n_mot)]
        nao_alocadas: dict[int, str] = {}
        if com_posicao:
            try:
                duracoes, _ = provedor.matriz([partida, *[paradas[i] for i in com_posicao]])
            except MatrizIndisponivel as exc:
                raise HTTPException(503, str(exc)) from None
            solucao = resolver(duracoes, n_mot, limite, retorno, tempo_para(len(com_posicao), cfg.solver_tempo_limite_s))
            rotas = [[com_posicao[i] for i in rota] for rota in solucao.rotas]
            for i, motivo in solucao.nao_alocadas.items():
                if motivo == MOTIVO_LIMITE:
                    nao_alocadas[com_posicao[i]] = motivo
                else:
                    sem_posicao.append(com_posicao[i])
        for i in sorted(sem_posicao):
            livres = [r for r in rotas if not limite or len(r) < limite]
            if livres:
                min(livres, key=len).append(i)
            else:
                nao_alocadas[i] = MOTIVO_LIMITE
        tempo_calculo_ms = int((time.monotonic() - inicio) * 1000)

        sem_mapa = set(sem_posicao)
        # No link, parada sem posição vai só pelo texto (nunca por uma coordenada errada)
        para_link = [Ponto(None, None, p.texto) if i in sem_mapa else p for i, p in enumerate(paradas)]
        motoristas = []
        for m, rota in zip(dados.motoristas, rotas):
            item = {"nome": m.nome, "paradas": rota, "distancia_m": 0.0, "duracao_s": 0.0, "linha": [], "trechos": []}
            if rota:
                no_mapa = [paradas[i] for i in rota if i not in sem_mapa]
                seq = [partida, *no_mapa] + ([partida] if retorno else [])
                if len(seq) > 1:
                    try:
                        item["distancia_m"], item["duracao_s"], item["linha"] = provedor.rota(seq)
                    except MatrizIndisponivel as exc:
                        raise HTTPException(503, str(exc)) from None
                item["trechos"] = trechos(partida, [para_link[i] for i in rota], retorno)
            motoristas.append(item)

    nao_alocadas = [{"indice": i, "motivo": motivo} for i, motivo in sorted(nao_alocadas.items())]
    km_total = sum(m["distancia_m"] for m in motoristas) / 1000
    tempo_total = sum(m["duracao_s"] for m in motoristas)
    db.add(MetricaUso(usuario_id=usuario.id, qtd_paradas=len(paradas), qtd_motoristas=len(dados.motoristas),
                      km_total=round(km_total, 2), tempo_total_s=int(tempo_total), tempo_calculo_ms=tempo_calculo_ms))
    db.commit()
    return {
        "partida": {"lat": partida.lat, "lng": partida.lng, "endereco": partida.texto},
        "fim": dados.fim,
        "motoristas": motoristas,
        "nao_alocadas": nao_alocadas,
        "sem_posicao": sorted(sem_posicao),  # nas rotas, mas sem ponto confiável no mapa (fora do km/tempo)
        "km_total": round(km_total, 2),
        "tempo_total_s": int(tempo_total),
        "tempo_calculo_ms": tempo_calculo_ms,
    }


# --- Exportação ---------------------------------------------------------------


class ParadaExport(BaseModel):
    endereco: str = Field(max_length=400)
    referencia: str = Field(default="", max_length=200)
    volumes: str = Field(default="", max_length=20)
    telefone: str = Field(default="", max_length=60)
    pedido: str = Field(default="", max_length=100)
    observacoes: str = Field(default="", max_length=300)
    lat: float | None = None
    lng: float | None = None


class TrechoExport(BaseModel):
    url: str = Field(max_length=2100)


class MotoristaExport(BaseModel):
    nome: str = Field(max_length=60)
    distancia_m: float
    duracao_s: float
    paradas: list[ParadaExport]
    trechos: list[TrechoExport]


class NaoAlocadaExport(BaseModel):
    endereco: str = Field(max_length=400)
    motivo: str = Field(max_length=200)


class ExportIn(BaseModel):
    motoristas: list[MotoristaExport] = Field(max_length=MAX_MOTORISTAS)
    nao_alocadas: list[NaoAlocadaExport] = []


@router.post("/api/exportar/xlsx")
def exportar_xlsx(dados: ExportIn, _: Usuario = Depends(exigir_login)):
    wb = Workbook()
    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="3056F5")

    def preparar(ws, titulos, larguras):
        ws.append(titulos)
        for i, (c, w) in enumerate(zip(ws[1], larguras)):
            c.font, c.fill = cabecalho, fundo
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Rotas"
    preparar(ws, ["Motorista", "Ordem", "Endereço", "Cliente / referência", "Telefone", "Pedido", "Volumes", "Observações",
                  "Latitude", "Longitude"], [18, 8, 60, 28, 18, 12, 10, 36, 12, 12])
    for m in dados.motoristas:
        for ordem, p in enumerate(m.paradas, start=1):
            ws.append([m.nome, ordem, p.endereco, p.referencia, p.telefone, p.pedido, p.volumes, p.observacoes,
                       round(p.lat, 6) if p.lat is not None else None, round(p.lng, 6) if p.lng is not None else None])

    resumo = wb.create_sheet("Resumo")
    preparar(resumo, ["Motorista", "Paradas", "Km", "Tempo estimado (min)"], [18, 10, 10, 22])
    for m in dados.motoristas:
        resumo.append([m.nome, len(m.paradas), round(m.distancia_m / 1000, 1), round(m.duracao_s / 60)])

    links = wb.create_sheet("Links")
    preparar(links, ["Motorista", "Trecho", "Link do Google Maps"], [18, 8, 100])
    for m in dados.motoristas:
        for i, t in enumerate(m.trechos, start=1):
            links.append([m.nome, i, t.url])
            c = links.cell(row=links.max_row, column=3)
            c.hyperlink, c.font = t.url, Font(color="1F3FD1", underline="single")
            c.alignment = Alignment(wrap_text=False)

    if dados.nao_alocadas:
        na = wb.create_sheet("Não alocadas")
        preparar(na, ["Endereço", "Motivo"], [60, 50])
        for p in dados.nao_alocadas:
            na.append([p.endereco, p.motivo])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"rotas_{datetime.now(FUSO):%Y-%m-%d_%H%M}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"', "Cache-Control": "no-store"},
    )
