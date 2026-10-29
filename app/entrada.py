"""Leitura de planilhas (.xlsx/.csv) só em memória: nada é gravado em disco ou no banco."""
import csv
import io
import re
import unicodedata
import zipfile

from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_DESCOMPACTADO = 60 * 1024 * 1024  # proteção contra "zip bomb" em .xlsx
MAX_LINHAS = 5000
MAX_COLUNAS = 60

# Campo -> nomes de coluna reconhecidos (comparados sem acento, minúsculos, só letras/números)
SINONIMOS = {
    "endereco": ["endereco", "enderecocompleto", "logradouro", "rua", "address", "end", "enderecoentrega", "local"],
    "numero": ["numero", "num", "no", "n", "nro", "numeroresidencia"],
    "complemento": ["complemento", "compl", "comp"],
    "bairro": ["bairro", "setor", "regiao", "ra"],
    "cidade": ["cidade", "municipio", "localidade"],
    "uf": ["uf", "estado"],
    "cep": ["cep", "codigopostal"],
    "referencia": ["nome", "referencia", "ref", "destinatario", "cliente", "pedido", "identificacao", "id"],
    "volumes": ["volumes", "volume", "qtd", "qtde", "quantidade", "vol", "caixas", "itens"],
}


class ArquivoInvalido(Exception):
    pass


def _chave(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


def sugerir_mapeamento(colunas: list[str]) -> dict[str, int | None]:
    chaves = [_chave(c) for c in colunas]
    sugestao: dict[str, int | None] = {}
    usados: set[int] = set()
    for campo, nomes in SINONIMOS.items():
        sugestao[campo] = None
        for nome in nomes:  # a ordem dos sinônimos define a prioridade
            for i, ch in enumerate(chaves):
                if i not in usados and ch == nome:
                    sugestao[campo] = i
                    usados.add(i)
                    break
            if sugestao[campo] is not None:
                break
    return sugestao


def _texto_celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # 70040.0 -> 70040 (CEP, número)
    return " ".join(str(v).split())


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    if not conteudo.startswith(b"PK\x03\x04"):
        raise ArquivoInvalido("O arquivo não é uma planilha .xlsx válida.")
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPACTADO:
                raise ArquivoInvalido("A planilha é grande demais depois de descompactada.")
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except ArquivoInvalido:
        raise
    except Exception as exc:
        raise ArquivoInvalido("Não foi possível ler a planilha .xlsx.") from exc
    try:
        ws = wb.worksheets[0]
        linhas = []
        for linha in ws.iter_rows(values_only=True, max_col=MAX_COLUNAS):
            linhas.append([_texto_celula(v) for v in linha])
            if len(linhas) > MAX_LINHAS + 1:
                raise ArquivoInvalido(f"A planilha tem mais de {MAX_LINHAS} linhas.")
        return linhas
    finally:
        wb.close()


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    if b"\x00" in conteudo[:4096]:
        raise ArquivoInvalido("O arquivo não parece ser um CSV de texto.")
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ArquivoInvalido("Não foi possível ler o texto do CSV.")
    amostra = texto[:8192]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
        delim = dialeto.delimiter
    except csv.Error:
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
    linhas = []
    for linha in csv.reader(io.StringIO(texto), delimiter=delim):
        linhas.append([_texto_celula(v) for v in linha[:MAX_COLUNAS]])
        if len(linhas) > MAX_LINHAS + 1:
            raise ArquivoInvalido(f"O CSV tem mais de {MAX_LINHAS} linhas.")
    return linhas


def ler_planilha(nome: str, conteudo: bytes) -> dict:
    if len(conteudo) > MAX_BYTES:
        raise ArquivoInvalido("O arquivo passa de 5 MB.")
    if not conteudo:
        raise ArquivoInvalido("O arquivo está vazio.")
    extensao = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    if extensao == "xlsx":
        linhas = _ler_xlsx(conteudo)
    elif extensao == "csv":
        linhas = _ler_csv(conteudo)
    else:
        raise ArquivoInvalido("Envie um arquivo .xlsx ou .csv.")

    linhas = [l for l in linhas if any(c for c in l)]
    if len(linhas) < 2:
        raise ArquivoInvalido("A planilha precisa de uma linha de cabeçalho e pelo menos uma linha de dados.")
    largura = max(len(l) for l in linhas)
    while largura > 0 and not any(len(l) >= largura and l[largura - 1] for l in linhas):
        largura -= 1  # remove colunas vazias à direita
    linhas = [(l + [""] * largura)[:largura] for l in linhas]
    colunas = [c or f"Coluna {i + 1}" for i, c in enumerate(linhas[0])]
    return {"colunas": colunas, "linhas": linhas[1:], "sugestao": sugerir_mapeamento(colunas)}
