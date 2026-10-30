import io

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Configuracao, MetricaUso
from tests.conftest import login

PARTIDA = {"lat": -15.80, "lng": -47.90}


def _paradas(n):
    return [{"lat": -15.75 - i * 0.003, "lng": -47.88 - (i % 5) * 0.004, "endereco": f"Endereço {i}"} for i in range(n)]


def _logar(cliente, criar_usuario):
    criar_usuario()
    login(cliente, "op@teste.com")


def test_calculo_completo(cliente, criar_usuario):
    _logar(cliente, criar_usuario)
    r = cliente.post("/api/calcular", json={"partida": PARTIDA, "paradas": _paradas(23), "motoristas": [{"nome": "João"}, {}]})
    assert r.status_code == 200, r.text
    d = r.json()
    assert [m["nome"] for m in d["motoristas"]] == ["João", "Motorista 2"]
    todas = sorted(i for m in d["motoristas"] for i in m["paradas"])
    assert todas == list(range(23)) and d["nao_alocadas"] == []
    for m in d["motoristas"]:
        assert m["distancia_m"] > 0 and m["linha"] and m["trechos"]
    with SessionLocal() as db:
        met = db.query(MetricaUso).one()
        assert met.qtd_paradas == 23 and met.qtd_motoristas == 2 and met.km_total > 0


def test_partida_obrigatoria_e_dentro_da_area(cliente, criar_usuario):
    _logar(cliente, criar_usuario)
    r = cliente.post("/api/calcular", json={"paradas": _paradas(2), "motoristas": [{}]})
    assert r.status_code == 422 and "localização" in r.json()["detail"]
    r = cliente.post("/api/calcular", json={"partida": {"lat": -23.55, "lng": -46.63}, "paradas": _paradas(2), "motoristas": [{}]})
    assert r.status_code == 422 and "fora da área" in r.json()["detail"]


def test_endereco_fixo_ignora_partida_do_operador(cliente, criar_usuario):
    with SessionLocal() as db:
        c = db.get(Configuracao, 1)
        c.modo_partida, c.partida_endereco, c.partida_lat, c.partida_lng = "endereco_fixo", "SIA Trecho 3", -15.81, -47.95
        db.commit()
    _logar(cliente, criar_usuario)
    r = cliente.post("/api/calcular", json={"partida": {"lat": -15.6, "lng": -47.7}, "paradas": _paradas(3), "motoristas": [{}]})
    assert r.json()["partida"] == {"lat": -15.81, "lng": -47.95, "endereco": "SIA Trecho 3"}


def test_parada_sem_posicao_entra_no_fim_da_rota(cliente, criar_usuario):
    """Não localizada ou fora da área: entra na rota (no fim) e vai para o Google Maps só pelo texto."""
    _logar(cliente, criar_usuario)
    paradas = _paradas(3) + [{"lat": -23.55, "lng": -46.63, "endereco": "Riacho Fundo I, AC 02 Lote 1"},
                             {"endereco": "QNM 34 Conjunto H Casa 5"}]
    d = cliente.post("/api/calcular", json={"partida": PARTIDA, "paradas": paradas, "motoristas": [{}]}).json()
    assert d["nao_alocadas"] == [] and d["sem_posicao"] == [3, 4]
    rota = d["motoristas"][0]["paradas"]
    assert sorted(rota) == [0, 1, 2, 3, 4] and rota[-2:] == [3, 4]
    url = d["motoristas"][0]["trechos"][-1]["url"]
    assert "Riacho%20Fundo%20I" in url and "-23.55" not in url


def test_tempo_do_otimizador_proporcional():
    from app.roteamento.solver import tempo_para
    assert tempo_para(5, 30) == 1 and tempo_para(15, 30) == 1
    assert tempo_para(40, 30) == 5 and tempo_para(41, 30) == 30
    assert tempo_para(30, 3) == 3  # nunca passa do limite do admin


def test_limites(cliente, criar_usuario):
    _logar(cliente, criar_usuario)
    with SessionLocal() as db:
        db.get(Configuracao, 1).max_paradas = 5
        db.commit()
    r = cliente.post("/api/calcular", json={"partida": PARTIDA, "paradas": _paradas(6), "motoristas": [{}]})
    assert r.status_code == 422
    r = cliente.post("/api/calcular", json={"partida": PARTIDA, "paradas": _paradas(5), "motoristas": [{}], "max_por_motorista": 2})
    assert len(r.json()["nao_alocadas"]) == 3


def test_exportar_xlsx(cliente, criar_usuario):
    _logar(cliente, criar_usuario)
    corpo = {
        "motoristas": [{"nome": "Ana", "distancia_m": 12000, "duracao_s": 1800,
                        "paradas": [{"endereco": "Rua 1", "referencia": "Loja", "volumes": "2",
                                     "telefone": "(61) 99999-0000", "pedido": "1001", "observacoes": "Portaria", "lat": -15.7, "lng": -47.9}],
                        "trechos": [{"url": "https://www.google.com/maps/dir/?api=1&origin=a&destination=b&travelmode=driving"}]}],
        "nao_alocadas": [{"endereco": "Rua X", "motivo": "Sem rota"}],
    }
    r = cliente.post("/api/exportar/xlsx", json=corpo)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Rotas", "Resumo", "Links", "Não alocadas"]
    assert [c.value for c in wb["Rotas"][2]][:8] == ["Ana", 1, "Rua 1", "Loja", "(61) 99999-0000", "1001", "2", "Portaria"]
    assert wb["Resumo"]["C2"].value == 12.0


def test_config_do_planejamento(cliente, criar_usuario):
    _logar(cliente, criar_usuario)
    c = cliente.get("/api/planejamento/config").json()
    assert c["modo_partida"] == "localizacao_atual" and c["partida"] is None and c["max_paradas"] == 200
