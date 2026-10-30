"""Gera app/static/modelos/modelo-entregas.xlsx: o modelo padrão que a operação baixa e preenche ao longo do dia."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# (título, largura, obrigatória, dica). Os nomes casam com os sinônimos de app/entrada.py.
COLUNAS = [
    ("Cliente", 28, False, "Nome de quem comprou."),
    ("Telefone", 18, False, "Com DDD. Ex.: (61) 99999-0000"),
    ("E-mail", 30, False, "E-mail do cliente."),
    ("Endereço completo", 60, True, "Tudo numa célula, como você escreveria no Google Maps. Ex.: Riacho Fundo I, AC 02 Lote 1"),
    ("Pedido", 14, False, "Número do pedido ou nota, se houver."),
    ("Volumes", 10, False, "Quantidade de caixas/pacotes."),
    ("Observações", 36, False, "Ponto de referência, horário, recado para o motorista."),
]
EXEMPLOS = [
    ["Maria (exemplo)", "(61) 99999-0000", "maria@exemplo.com", "SQS 308 Bloco C Apto 101, Asa Sul, Brasília - DF",
     "1001", 2, "Portaria 24h"],
    ["João (exemplo)", "(61) 98888-0000", "", "Riacho Fundo I, AC 02 Lote 1", "1002", 1, "Casa de muro azul"],
]
AZUL, CINZA = "3056F5", "F1F3F9"
pasta = Path(__file__).resolve().parent.parent / "app" / "static" / "modelos"
pasta.mkdir(parents=True, exist_ok=True)


def cabecalho(ws, colunas):
    ws.append([c[0] + (" *" if c[2] else "") for c in colunas])
    for i, (c, (_, largura, _, dica)) in enumerate(zip(ws[1], colunas)):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center")
        c.comment = Comment(dica, "Rotear")
        ws.column_dimensions[c.column_letter].width = largura
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


wb = Workbook()

# 1ª aba: é a que o sistema lê. Fica só com o cabeçalho para não importar exemplos por engano.
ws = wb.active
ws.title = "Entregas"
cabecalho(ws, COLUNAS)
letra = {c[0]: ws.cell(row=1, column=i + 1).column_letter for i, c in enumerate(COLUNAS)}
for nome in ("Telefone", "Pedido"):  # texto: preserva zeros e não vira notação científica
    for linha in range(2, 1001):
        ws[f"{letra[nome]}{linha}"].number_format = "@"
volumes = DataValidation(type="whole", operator="between", formula1="1", formula2="999", allow_blank=True,
                         showErrorMessage=True, errorTitle="Volumes", error="Informe um número inteiro.")
ws.add_data_validation(volumes)
volumes.add(f"{letra['Volumes']}2:{letra['Volumes']}1000")

# 2ª aba: como preencher
ins = wb.create_sheet("Como preencher")
ins.column_dimensions["A"].width = 100
textos = [
    ("Como usar este modelo", True),
    ("1. Ao longo do dia, anote cada venda em uma linha da aba \"Entregas\" (a primeira aba).", False),
    ("2. No fim do dia, salve o arquivo e importe no Rotear: passo \"Endereços\" > \"Planilha\".", False),
    ("3. Não mude o nome nem a ordem das colunas. Não deixe linhas em branco no meio.", False),
    ("4. Só o Endereço completo (*) é obrigatório. Escreva tudo numa célula, do jeito que você pesquisaria no Google Maps.", False),
    ("5. Incluir a cidade ou região (ex.: Riacho Fundo I, Ceilândia) ajuda a montar a melhor ordem de entregas.", False),
    ("6. Use um arquivo novo por dia (ex.: entregas-2026-10-30.xlsx). Salve sempre como .xlsx.", False),
    ("", False),
    ("Exemplo de preenchimento (não copie estas linhas para a aba Entregas):", True),
]
for texto, negrito in textos:
    ins.append([texto])
    ins.cell(row=ins.max_row, column=1).font = Font(bold=negrito, size=13 if negrito else 11)
ins.append([])
inicio = ins.max_row + 1
ins.append([c[0] for c in COLUNAS])
for linha in EXEMPLOS:
    ins.append(linha)
for c in ins[inicio]:
    c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=AZUL)
for r in ins.iter_rows(min_row=inicio + 1, max_row=ins.max_row):
    for c in r:
        c.fill = PatternFill("solid", fgColor=CINZA)
        c.font = Font(italic=True)

wb.save(pasta / "modelo-entregas.xlsx")
print(pasta / "modelo-entregas.xlsx")
